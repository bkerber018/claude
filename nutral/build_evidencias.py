"""Gera Evidencias_Nutral_ML_2026-10-20.xlsx a partir das consultas JoomPulse de 20/10/2026.

Os dados abaixo são cópia literal das respostas do JoomPulse (snapshot de produtos com
~24h de lag; categoria mensal com ~31 dias de lag). Vendas e GMV são estimativas JoomPulse.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CAPTURA = "20/10/2026"
FONTE_PROD = "JoomPulse · MercadoProductsWeekly (snapshot ~24h)"
FONTE_CAT = "JoomPulse · MlbCategoriesMonthly"

# productId, título, vendedor, marca cadastrada, loja oficial, buy box, preço, pedidos 1m (est.), GMV 1m (est.)
OFERTAS = [
    ("MLB19939966", "Ração Nutra Fish 1kg Em Flocos Comida Para Peixes", "LS20260722105541078", "NUTRAL", False, False, 148.99, 0, 0),
    ("MLB19939966", "Ração Nutra Fish 1kg Em Flocos Comida Para Peixes", "TURMA PET", "NUTRAL", False, False, 149.99, 0, 0),
    ("MLB19939966", "Ração Nutra Fish 1kg Em Flocos Comida Para Peixes", "AQUÁRIOSTOP SHOW", "NUTRAL", False, False, 173.85, 0, 0),
    ("MLB19939966", "Ração Nutra Fish 1kg Em Flocos Comida Para Peixes", "AQUARIUM_SHOP", "NUTRAL", False, False, 179.99, 0, 0),
    ("MLB19939966", "Ração Nutra Fish 1kg Em Flocos Comida Para Peixes", "AQUATURBO_SHOP", "NUTRAL", False, False, 189.90, 0, 0),
    ("MLB19939966", "Ração Nutra Fish 1kg Em Flocos Comida Para Peixes", "REIDOPETS", "NUTRAL", False, False, 199.00, 0, 0),
    ("MLB19939966", "Ração Nutra Fish 1kg Em Flocos Comida Para Peixes", "ACQUAPEX", "NUTRAL", True, True, 210.50, 63, 13261.50),
    ("MLB19939966", "Ração Nutra Fish 1kg Em Flocos Comida Para Peixes", "ACQUAPEX", "NUTRAL", True, False, 224.82, 0, 0),
    ("MLB19984593", "Ração P/tartarugas E Cágados Criador Nutral Ouro 1,3kg", "NETE4UVARIEDADES", "NUTRAL", True, False, 178.00, 0, 0),
    ("MLB19984593", "Ração P/tartarugas E Cágados Criador Nutral Ouro 1,3kg", "LS20260722105541078", "NUTRAL", False, False, 188.99, 0, 0),
    ("MLB19984593", "Ração P/tartarugas E Cágados Criador Nutral Ouro 1,3kg", "TURMA PET", "NUTRAL", False, True, 189.99, 46, 8739.54),
    ("MLB19984593", "Ração P/tartarugas E Cágados Criador Nutral Ouro 1,3kg", "NETE4UVARIEDADES", "NUTRAL", True, False, 193.99, 0, 0),
    ("MLB19984593", "Ração P/tartarugas E Cágados Criador Nutral Ouro 1,3kg", "FISCHCAT DOG", "NUTRAL", False, False, 236.55, 0, 0),
    ("MLB19984593", "Ração P/tartarugas E Cágados Criador Nutral Ouro 1,3kg", "AQUATURBO_SHOP", "NUTRAL", False, False, 236.90, 0, 0),
    ("MLB19999688", "Ração PARA Carpas Nutral 2kg", "NETE4UVARIEDADES", "NUTRAL", True, False, 178.10, 0, 0),
    ("MLB19999688", "Ração PARA Carpas Nutral 2kg", "LS20260722105541078", "NUTRAL", False, False, 185.24, 0, 0),
    ("MLB19999688", "Ração PARA Carpas Nutral 2kg", "SILVACIAVARIEDADES", "NUTRAL", True, False, 187.99, 0, 0),
    ("MLB19999688", "Ração PARA Carpas Nutral 2kg", "CARMO_PETS", "NUTRAL", False, False, 219.99, 0, 0),
    ("MLB19999688", "Ração PARA Carpas Nutral 2kg", "AQUÁRIOSTOP SHOW", "NUTRAL", False, True, 228.00, 15, 3420.00),
    ("MLB19999688", "Ração PARA Carpas Nutral 2kg", "FISCHCAT DOG", "NUTRAL", False, False, 245.00, 0, 0),
    ("MLB19999688", "Ração PARA Carpas Nutral 2kg", "YELLOWFISH ECOMMERCE", "NUTRAL", False, False, 252.00, 0, 0),
    ("MLB20033085", "Ração Nutral Ouro Tartaruga - 300 G", "LS20260722105541078", "NUTRAL", False, False, 75.89, 0, 0),
    ("MLB20033085", "Ração Nutral Ouro Tartaruga - 300 G", "TURMA PET", "NUTRAL", False, False, 77.99, 0, 0),
    ("MLB20033085", "Ração Nutral Ouro Tartaruga - 300 G", "NETE4UVARIEDADES", "NUTRAL", True, True, 78.00, 21, 1638.00),
    ("MLB21461130", "Suplemento Aminoácidos para Cães Nutral 300g", "ZEDANET", "NUTRAL", False, False, 61.73, 0, 0),
    ("MLB21461130", "Suplemento Aminoácidos para Cães Nutral 300g", "CARMO_PETS", "NUTRAL", False, True, 64.99, 65, 4224.35),
    ("MLB21461130", "Suplemento Aminoácidos para Cães Nutral 300g", "BESTFREND01", "NUTRAL", False, False, 65.90, 0, 0),
    ("MLB21461130", "Suplemento Aminoácidos para Cães Nutral 300g", "TURMA PET", "NUTRAL", False, False, 72.49, 0, 0),
    ("MLB21461130", "Suplemento Aminoácidos para Cães Nutral 300g", "GATOGORDO", "NUTRAL", False, False, 74.80, 0, 0),
    ("MLB21461130", "Suplemento Aminoácidos para Cães Nutral 300g", "BAKANUTILIDADES", "NUTRAL", False, False, 76.90, 0, 0),
    ("MLB19946397", "Ração Mega Food Jabuti 200g Com Alho", "FISCHCAT DOG", "MEGAFOOD", False, False, 49.00, 0, 0),
    ("MLB19946397", "Ração Mega Food Jabuti 200g Com Alho", "NETE4UVARIEDADES", "MEGAFOOD", True, False, 50.72, 0, 0),
    ("MLB19946397", "Ração Mega Food Jabuti 200g Com Alho", "LS20260722105541078", "MEGAFOOD", False, False, 53.99, 0, 0),
    ("MLB19946397", "Ração Mega Food Jabuti 200g Com Alho", "TURMA PET", "MEGAFOOD", False, False, 54.99, 0, 0),
    ("MLB19946397", "Ração Mega Food Jabuti 200g Com Alho", "AQUATURBO_SHOP", "MEGAFOOD", False, True, 58.90, 57, 3357.30),
    ("MLB19946397", "Ração Mega Food Jabuti 200g Com Alho", "REIDOPETS", "MEGAFOOD", False, False, 64.90, 0, 0),
    ("MLB19946397", "Ração Mega Food Jabuti 200g Com Alho", "AQUÁRIOSTOP SHOW", "MEGAFOOD", False, False, 65.00, 0, 0),
    ("MLB19946397", "Ração Mega Food Jabuti 200g Com Alho", "ACQUAPEX", "MEGAFOOD", True, False, 73.36, 0, 0),
    ("MLB19946397", "Ração Mega Food Jabuti 200g Com Alho", "YELLOWFISH ECOMMERCE", "MEGAFOOD", False, False, 139.01, 0, 0),
    ("MLB19946162", "Ração Para Jabuti Mega Food", "NETE4UVARIEDADES", "MEGAFOOD", True, False, 172.50, 0, 0),
    ("MLB19946162", "Ração Para Jabuti Mega Food", "LS20260722105541078", "MEGAFOOD", False, False, 173.99, 0, 0),
    ("MLB19946162", "Ração Para Jabuti Mega Food", "AQUÁRIOSTOP SHOW", "MEGAFOOD", False, False, 185.25, 0, 0),
    ("MLB19946162", "Ração Para Jabuti Mega Food", "AQUATURBO_SHOP", "MEGAFOOD", False, True, 198.90, 14, 2784.60),
    ("MLB19946162", "Ração Para Jabuti Mega Food", "FISCHCAT DOG", "MEGAFOOD", False, False, 202.00, 0, 0),
    ("MLB19938328", "Ração Mega Food Basica Bits Alho Media 500g Carpa Kinguio", "NETE4UVARIEDADES", "MEGAFOOD", True, False, 67.99, 0, 0),
    ("MLB19938328", "Ração Mega Food Basica Bits Alho Media 500g Carpa Kinguio", "TURMA PET", "MEGAFOOD", False, False, 72.99, 0, 0),
    ("MLB19938328", "Ração Mega Food Basica Bits Alho Media 500g Carpa Kinguio", "LS20260722105541078", "MEGAFOOD", False, False, 72.99, 0, 0),
    ("MLB19938328", "Ração Mega Food Basica Bits Alho Media 500g Carpa Kinguio", "AQUATURBO_SHOP", "MEGAFOOD", False, False, 78.90, 0, 0),
    ("MLB19938328", "Ração Mega Food Basica Bits Alho Media 500g Carpa Kinguio", "CLEOPATRALEONARDA", "MEGAFOOD", False, False, 79.90, 0, 0),
    ("MLB19938328", "Ração Mega Food Basica Bits Alho Media 500g Carpa Kinguio", "AQUARIUM_SHOP", "MEGAFOOD", False, False, 99.90, 0, 0),
    ("MLB19938328", "Ração Mega Food Basica Bits Alho Media 500g Carpa Kinguio", "YELLOWFISH ECOMMERCE", "MEGAFOOD", False, True, 142.55, 15, 2138.25),
    ("MLB19946628", "Ração Mega Food Carnivoros De Superfície 1,5kg Com Alho", "AQUÁRIOSTOP SHOW", "NUTRAL", False, False, 189.00, 0, 0),
    ("MLB19946628", "Ração Mega Food Carnivoros De Superfície 1,5kg Com Alho", "FISCHCAT DOG", "NUTRAL", False, True, 195.00, 14, 2730.00),
    ("MLB19946628", "Ração Mega Food Carnivoros De Superfície 1,5kg Com Alho", "YELLOWFISH ECOMMERCE", "NUTRAL", False, False, 285.64, 0, 0),
    ("MLB19946156", "Ração Tartaruga Nutricon Turtle 270g + Micro Camarão Aviu 60g Mega Food", "YELLOWFISH ECOMMERCE", "NUTRICONMEGAFOOD", False, True, 277.49, 10, 2774.90),
]

# marca, L2, GMV 1m est., pedidos 1m est., anúncios, vendedores, produtos — filtro L1 = Pet Shop
FAMILIA = [
    ("NUTRAL", "Peixes", 39278.38, 266, 148, 25, 84),
    ("MEGAFOOD", "Peixes", 12028.28, 143, 135, 14, 75),
    ("NUTRAL", "Cães", 4437.23, 68, 12, 6, 4),
    ("NUTRICONMEGAFOOD", "Peixes", 3061.82, 14, 5, 3, 4),
    ("NUTRAL", "Anfíbios e Répteis", 2482.66, 14, 11, 7, 11),
    ("NUTRALOURO", "Anfíbios e Répteis", 2095.21, 29, 5, 3, 5),
    ("MEGAFOOD", "Anfíbios e Répteis", 1971.00, 26, 11, 9, 11),
    ("MEGAFOOD", "Aves e Acessórios", 1723.50, 27, 3, 1, 2),
    ("NUTRAFISH", "Peixes", 1234.63, 15, 30, 11, 17),
]

# marca, GMV 1m est., pedidos 1m est., vendedores, produtos — categoria Alimentos (Peixes) MLB1094
MARCAS_CATEGORIA = [
    ("POYTARA", 758141.85, 5583, 319, 3731), ("TROPICAL", 249289.74, 1483, 171, 4732),
    ("NUTRICON", 180521.20, 1729, 149, 2172), ("POLINUTRI", 150603.88, 933, 41, 227),
    ("TETRA", 138888.45, 1009, 129, 808), ("ALCON", 133111.37, 1551, 233, 3631),
    ("MARAMAR", 69827.22, 928, 52, 1797), ("OCEANTECH", 60834.07, 959, 159, 1912),
    ("NUTRAL", 39278.38, 266, 25, 84), ("PRODACINTERNATIONAL", 23762.10, 161, 70, 158),
    ("POLINUTRICHRONOS", 16842.82, 329, 11, 36), ("GUABI", 16719.77, 399, 20, 94),
    ("REDSEA", 16035.21, 57, 39, 191), ("VITALIS", 15227.28, 108, 44, 254),
    ("PRODAC", 14904.67, 37, 55, 1507), ("POYTARARACOES", 13582.99, 48, 26, 15),
    ("LAGUNA", 13493.45, 159, 18, 73), ("MEGAFOOD", 12028.28, 143, 14, 75),
]

# mês, GMV est., pedidos est., vendedores, marcas, participação do maior vendedor
CATEGORIA_MENSAL = [
    ("2025-09", 3288157.64, 34059, 1009, 816, 0.1352), ("2025-10", 3594774.50, 36483, 1054, 826, 0.1283),
    ("2025-11", 3829066.02, 36013, 1087, 841, 0.1364), ("2025-12", 3873641.45, 33232, 1081, 852, 0.1473),
    ("2026-01", 4250878.97, 38087, 944, 862, 0.1427), ("2026-02", 3790014.40, 35692, 971, 857, 0.1221),
    ("2026-03", 4756690.72, 40163, 1015, 847, 0.1534), ("2026-04", 3705302.06, 39003, 982, 802, 0.1501),
    ("2026-05", 3549716.20, 38335, 970, 826, 0.1539), ("2026-06", 3911045.19, 35776, 968, 783, 0.1410),
    ("2026-07", 3366427.40, 35694, 1037, 787, 0.1357), ("2026-08", 2698686.84, 31073, 1089, 792, 0.1458),
    ("2026-09", 1478240.99, 16237, 1171, 774, 0.1553),
]

MUDANCAS = [
    ("Carpa 2 kg · MLB19999688", "Buy box", "POPi BOX EMBALAGENS (deck, slide 5)", "AQUÁRIOSTOP SHOW a R$ 228,00", "JoomPulse"),
    ("Carpa 2 kg · MLB19999688", "Faixa de preço", "R$ 178,10 a R$ 299,90", "R$ 178,10 a R$ 252,00 (7 ofertas)", "Print 20/10 + JoomPulse"),
    ("Carpa 2 kg · MLB19999688", "POPi BOX na lista de ofertas", "1 oferta citada", "2 ofertas com o nome POPi BOX EMBALAGENS (R$ 178,10 · +5 mil vendas e R$ 191,95 · +1000 vendas); JoomPulse liga a 2 contas: NETE4UVARIEDADES e SILVACIAVARIEDADES", "Print 20/10 + JoomPulse"),
    ("Tartaruga Criador 1,3 kg · MLB19984593", "Buy box", "POPi BOX EMBALAGENS", "TURMA PET a R$ 189,99", "JoomPulse"),
    ("Tartaruga Criador 1,3 kg · MLB19984593", "Faixa de preço", "R$ 184,29 a R$ 346,90", "R$ 178,00 a R$ 236,90 (6 anúncios)", "JoomPulse"),
    ("Nutrafish Básica 1 kg · MLB19939966", "ID do catálogo", "Não identificado (planilha)", "MLB19939966 — produto de maior venda estimada da família", "JoomPulse"),
    ("Mega Food (mesma indústria)", "Escopo", "Excluída como ruído (planilha, Nutrafish Bits e Ouro Beta)", "MEGAFOOD: 14 vendedores, 75 produtos, GMV est. R$ 12.028,28/mês em Alimentos (Peixes)", "JoomPulse + catálogo 2021"),
]

HEADER_FILL = PatternFill("solid", fgColor="0B1620")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def sheet(wb, title, subtitle, header, rows, widths, money_cols=(), pct_cols=()):
    ws = wb.create_sheet(title)
    ws.append([title.upper()])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([subtitle])
    ws["A2"].font = Font(italic=True, color="555555")
    ws.append(header)
    for c in ws[3]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=4):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if c.column in money_cols and isinstance(c.value, (int, float)):
                c.number_format = '"R$" #,##0.00'
            if c.column in pct_cols and isinstance(c.value, (int, float)):
                c.number_format = "0.0%"
    ws.freeze_panes = "A4"
    return ws


def main(out):
    wb = Workbook()
    wb.remove(wb.active)

    # Buy box x menor preço, por catálogo
    resumo = []
    for pid in dict.fromkeys(o[0] for o in OFERTAS):
        ofs = [o for o in OFERTAS if o[0] == pid]
        win = next(o for o in ofs if o[5])
        low = min(ofs, key=lambda o: o[6])
        resumo.append((pid, win[1], len(ofs), len({o[2] for o in ofs}), low[6], low[2],
                       win[6], win[2], win[6] / low[6] - 1, max(o[6] for o in ofs), win[7], win[8]))
    sheet(wb, "Buy box x menor preço",
          f"Captura {CAPTURA} · {FONTE_PROD} · buy box é retrato do momento · pedidos e GMV são estimativas JoomPulse",
          ["Catálogo", "Produto", "Anúncios", "Vendedores", "Menor preço", "Vendedor do menor preço",
           "Preço do buy box", "Vendedor no buy box", "Buy box acima do menor", "Maior preço",
           "Pedidos 1m (est.)", "GMV 1m (est.)"],
          resumo, [14, 44, 9, 10, 12, 22, 12, 22, 12, 12, 11, 13],
          money_cols=(5, 7, 10, 12), pct_cols=(9,))

    sheet(wb, "Ofertas por catálogo",
          f"Captura {CAPTURA} · {FONTE_PROD} · nome do vendedor é o cadastro; o nome de vitrine pode diferir (ex.: POPi BOX)",
          ["Catálogo", "Título", "Vendedor (cadastro)", "Marca cadastrada", "Loja oficial", "Buy box",
           "Preço", "Pedidos 1m (est.)", "GMV 1m (est.)"],
          [(*o[:4], "Sim" if o[4] else "Não", "Sim" if o[5] else "Não", *o[6:]) for o in OFERTAS],
          [14, 48, 24, 18, 8, 8, 11, 11, 13], money_cols=(7, 9))

    sheet(wb, "Família Nutral",
          f"Captura {CAPTURA} · {FONTE_PROD} · Pet Shop, anúncios ativos · marcas da Nutral Indústria como cadastradas",
          ["Marca cadastrada", "Subcategoria", "GMV 1m (est.)", "Pedidos 1m (est.)", "Anúncios", "Vendedores", "Produtos"],
          FAMILIA, [20, 22, 14, 12, 10, 11, 10], money_cols=(3,))

    sheet(wb, "Marcas em Alimentos",
          f"Captura {CAPTURA} · {FONTE_PROD} · categoria Pet Shop > Peixes > Alimentos (MLB1094) · top 18 marcas por GMV estimado",
          ["Marca", "GMV 1m (est.)", "Pedidos 1m (est.)", "Vendedores", "Produtos"],
          MARCAS_CATEGORIA, [22, 15, 13, 11, 10], money_cols=(2,))

    sheet(wb, "Categoria mensal",
          f"{FONTE_CAT} · MLB1094 · estimativa por velocidade de avaliações · set/2026 parece incompleto: não usar",
          ["Mês", "GMV (est.)", "Pedidos (est.)", "Vendedores", "Marcas", "Participação do maior vendedor"],
          CATEGORIA_MENSAL, [10, 15, 13, 11, 9, 14], money_cols=(2,), pct_cols=(6,))

    sheet(wb, "Mudou desde 29-09",
          "Comparação entre o Raio-X/planilha de 29/09/2026 e os dados de 20/10/2026",
          ["Objeto", "Campo", "Em 29/09", "Em 20/10", "Fonte"],
          MUDANCAS, [30, 18, 34, 52, 20])

    sheet(wb, "Método e limites", "O que cada número prova e o que não prova",
          ["Tema", "Limite"],
          [("Vendas e GMV", "Estimativas JoomPulse sobre histórico de anúncios. Não são faturamento e não batem com o número da Nutral. Servem para comparar marcas e catálogos entre si."),
           ("Buy box", "Retrato do snapshot (~24h). Troca de vencedor é frequente; refazer print na véspera da reunião."),
           ("Categoria mensal", "Estimativa por velocidade de avaliações, ~31 dias de atraso. Setembro/2026 caiu pela metade de um mês para o outro, o que indica dado incompleto; usar agosto."),
           ("Nome do vendedor", "O JoomPulse traz o nome de cadastro. A vitrine pode exibir outro nome. A ligação NETE4UVARIEDADES/SILVACIAVARIEDADES ↔ POPi BOX EMBALAGENS vem do cruzamento de preço e selo com o print, e precisa de confirmação."),
           ("Loja oficial", "Indica que a conta tem status de loja oficial de alguma marca, não necessariamente da Nutral."),
           ("Marca cadastrada", "Agrupamento pelo campo marca do anúncio. Produto Nutral cadastrado com outra marca fica fora; produto Mega Food cadastrado como NUTRAL fica dentro.")],
          [20, 110])
    wb.save(out)


if __name__ == "__main__":
    import sys
    main(sys.argv[1])
