"""Gera Projecao_PetToys.xlsx com fórmulas vivas (mesma lógica de motor.py)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.workbook.defined_name import DefinedName
from build import BASE, FASE1, HORIZ, MES_INI, ANO_INI, MAT_PADRAO, L
from motor import SAZ_PET, MES

wb = Workbook()
AZ = PatternFill("solid", fgColor="111A2C"); IN = PatternFill("solid", fgColor="FFF4E5")
H = Font(bold=True, color="FFFFFF"); B = Font(bold=True); MUTE = Font(color="78839A", italic=True)
fino = Border(bottom=Side(style="thin", color="DDDDDD"))
BRL = '"R$" #,##0;[Red]-"R$" #,##0'; PCT = "0.0%"

# ── Premissas
ws = wb.active; ws.title = "Premissas"
ws["A1"] = "Pet Toys · Projeção da venda direta (Fase 2) · premissas editáveis"; ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = "Células em laranja são editáveis. Simulação, não meta de faturamento. impulsio · outubro 2026"; ws["A2"].font = MUTE
linhas = [
  ("Contrato impulsio", None, None, None),
  ("Setup total", BASE["setup"], BRL, "Setup", "Pago em parcelas iguais"),
  ("Setup · parcelas", BASE["parc"], "0", "Parc", ""),
  ("Mensalidade nos primeiros meses", 3500, BRL, "MensA", "Somada à parcela do setup dá R$ 6.000"),
  ("Primeiros meses (até o mês)", 6, "0", "MesesA", ""),
  ("Mensalidade depois (ML 4.500 + Amazon 1.700)", 6200, BRL, "MensB", ""),
  ("Variável sobre a venda direta", BASE["varP"], PCT, "VarP", "Só sobre o excedente do gatilho"),
  ("Gatilho do variável (venda direta/mês)", BASE["gat"], BRL, "Gat", "Abaixo disso, variável zero"),
  ("Operação Pet Toys", None, None, None),
  ("Potencial da venda direta em regime (R$/mês)", BASE["pot"], BRL, "Pot", "1,2% da prateleira de ~R$ 16 mi/mês (JoomPulse, set/26)"),
  ("Margem sobre a venda (líq. comissão, frete, imposto)", BASE["marg"], PCT, "Marg", "Faixa informada: 15% a 20%"),
  ("Mídia paga (% do GMV)", BASE["adsP"], PCT, "AdsP", "Paga à plataforma, fora do contrato"),
  ("Mídia começa no mês", BASE["adsIni"], "0", "AdsIni", "3º mês de venda direta"),
  ("Verba mínima de mídia (R$/mês)", BASE["piso"], BRL, "Piso", ""),
  ("Verba mínima até o mês", BASE["pisoAte"], "0", "PisoAte", ""),
]
r = 4
for item in linhas:
    if item[1] is None:
        ws.cell(r, 1, item[0]).font = B; r += 1; continue
    nome, v, fmt, nm, nota = item
    ws.cell(r, 1, nome); c = ws.cell(r, 2, v); c.number_format = fmt; c.fill = IN
    ws.cell(r, 3, nota).font = MUTE
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"Premissas!$B${r}")
    r += 1
r += 1
ws.cell(r, 1, "Sazonalidade (varejo pet) · índice do mês").font = B; r += 1
saz0 = r
for i, m in enumerate(MES):
    ws.cell(r + i, 1, m); c = ws.cell(r + i, 2, SAZ_PET[i]); c.fill = IN; c.number_format = "0.00"
wb.defined_names["Saz"] = DefinedName("Saz", attr_text=f"Premissas!$B${saz0}:$B${saz0+11}")
ws.column_dimensions["A"].width = 52; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 58

# ── Projeção
pj = wb.create_sheet("Projeção")
cab = ["Mês", "Período", "Mês do ano", "Fase", "Maturação", "Sazonalidade", "Venda direta (GMV)",
       "Margem", "Mídia", "Setup", "Mensalidade", "Variável", "Total impulsio", "Resultado", "Acumulado"]
for j, t in enumerate(cab, 1):
    c = pj.cell(1, j, t); c.font = H; c.fill = AZ; c.alignment = Alignment(horizontal="center", wrap_text=True)
mat = [0] * FASE1 + [x / 100 for x in MAT_PADRAO[1:]]
for k in range(HORIZ):
    i = k + 2; n = k + 1; idx = (MES_INI + k) % 12; ano = ANO_INI + (MES_INI + k) // 12
    pj.cell(i, 1, n); pj.cell(i, 2, f"{MES[idx]}/{str(ano)[2:]}"); pj.cell(i, 3, idx + 1)
    pj.cell(i, 4, "1 · Protagonista" if n <= FASE1 else "2 · + venda direta")
    c = pj.cell(i, 5, mat[k] if k < len(mat) else 1); c.fill = IN; c.number_format = "0%"
    pj.cell(i, 6, f"=INDEX(Saz,C{i})").number_format = "0.00"
    pj.cell(i, 7, f"=ROUND(Pot*E{i}*F{i}/500,0)*500")
    pj.cell(i, 8, f"=G{i}*Marg")
    pj.cell(i, 9, f"=IF(OR(A{i}<AdsIni,G{i}=0),0,IF(A{i}<=PisoAte,MAX(G{i}*AdsP,Piso),G{i}*AdsP))")
    pj.cell(i, 10, f"=IF(A{i}<=Parc,Setup/Parc,0)")
    pj.cell(i, 11, f"=IF(A{i}<=MesesA,MensA,MensB)")
    pj.cell(i, 12, f"=IF(G{i}>Gat,(G{i}-Gat)*VarP,0)")
    pj.cell(i, 13, f"=J{i}+K{i}+L{i}")
    pj.cell(i, 14, f"=H{i}-I{i}-M{i}")
    pj.cell(i, 15, f"=N{i}" if n == 1 else f"=O{i-1}+N{i}")
    for j in range(7, 16): pj.cell(i, j).number_format = BRL
    for j in range(1, 16): pj.cell(i, j).border = fino
last = HORIZ + 1
t = last + 2
pj.cell(t, 1, "Pior momento do caixa").font = B; pj.cell(t, 7, f"=MIN(O2:O{last})").number_format = BRL
pj.cell(t + 1, 1, "Acumulado no fim do horizonte").font = B; pj.cell(t + 1, 7, f"=O{last}").number_format = BRL
pj.cell(t + 2, 1, "Primeiro mês com acumulado positivo").font = B
pj.cell(t + 2, 7, f'=IFERROR(INDEX(B2:B{last},MATCH(TRUE,INDEX(O2:O{last}>0,0),0)),"fora do horizonte")')
pj.cell(t + 4, 1, "A maturação (coluna E) é editável: os 3 primeiros meses são Fase 1, sem venda direta.").font = MUTE
pj.cell(t + 5, 1, "A venda da rede de revendedores e eventuais pedidos 1P da Amazon não entram nesta conta.").font = MUTE
for j, w in enumerate([6, 9, 8, 18, 10, 11, 16, 13, 12, 11, 12, 11, 13, 13, 14], 1):
    pj.column_dimensions[pj.cell(1, j).column_letter].width = w
pj.freeze_panes = "C2"

lc = LineChart(); lc.title = "Resultado acumulado da venda direta"; lc.height = 9; lc.width = 22
lc.add_data(Reference(pj, min_col=15, min_row=1, max_row=last), titles_from_data=True)
lc.set_categories(Reference(pj, min_col=2, min_row=2, max_row=last)); lc.y_axis.numFmt = '#,##0'
pj.add_chart(lc, f"B{t+8}")
bc = BarChart(); bc.title = "Venda direta mensal (GMV)"; bc.height = 9; bc.width = 22
bc.add_data(Reference(pj, min_col=7, min_row=1, max_row=last), titles_from_data=True)
bc.set_categories(Reference(pj, min_col=2, min_row=2, max_row=last)); bc.y_axis.numFmt = '#,##0'
pj.add_chart(bc, f"J{t+8}")

wb.save("Projecao_PetToys.xlsx"); print("ok")
